import os
import uuid
from pathlib import Path

from pypdf import PdfReader
from pptx import Presentation
import fitz

from .vision import image_bytes_to_text


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".pptx",
    ".md",
    ".txt",
}


def clean_text(text: str) -> str:
    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line]
    return "\n".join(lines)


def split_text(text: str, chunk_size: int = 1000, overlap: int = 150):
    text = clean_text(text)

    if not text:
        return []

    chunks = []
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))

        chunk = text[start:end]

        if chunk.strip():
            chunks.append(chunk.strip())

        if end >= len(text):
            break

        start = end - overlap

    return chunks


def process_pdf(file_path: str):
    """
    Process both normal and scanned PDF pages.

    Normal pages:
        pypdf extracts the text.

    Scanned / poor-text pages:
        Render the original page as an image and send it
        to Gemini Vision for OCR.
    """

    reader = PdfReader(file_path)
    pdf_document = fitz.open(file_path)

    chunks = []

    for page_number, page in enumerate(reader.pages, start=1):

        extracted_text = page.extract_text() or ""
        cleaned = clean_text(extracted_text)

        # If the extracted text is too short, treat the page
        # as scanned / image-based and use Gemini Vision.
        if len(cleaned) < 150:

            print(
                f"Page {page_number}: weak PDF text detected. "
                f"Using Gemini Vision OCR..."
            )

            pdf_page = pdf_document[page_number - 1]

            pixmap = pdf_page.get_pixmap(
                matrix=fitz.Matrix(2, 2),
                alpha=False
            )

            image_bytes = pixmap.tobytes("png")

            ocr_text = image_bytes_to_text(
                image_bytes=image_bytes,
                mime_type="image/png"
            )

            page_text = ocr_text

            source_type = "scanned_pdf"

        else:
            print(
                f"Page {page_number}: normal PDF text extraction."
            )

            page_text = cleaned
            source_type = "pdf"

        page_chunks = split_text(page_text)

        for chunk in page_chunks:

            chunks.append({
                "id": str(uuid.uuid4()),
                "text": chunk,
                "source": os.path.basename(file_path),
                "page": page_number,
                "section": None,
                "source_type": source_type,
            })

    pdf_document.close()

    return chunks


def process_pptx(file_path: str):

    file_path = os.path.abspath(file_path)

    print(f"Opening PPTX: {file_path}")

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"PPTX file was not found at: {file_path}"
        )

    presentation = Presentation(file_path)

    chunks = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        texts = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    texts.append(text)

        slide_text = "\n".join(texts)

        if not slide_text:
            slide_text = "[No readable text found on this slide.]"

        for chunk in split_text(slide_text):

            chunks.append({
                "id": str(uuid.uuid4()),
                "text": chunk,
                "source": os.path.basename(file_path),
                "page": slide_number,
                "section": f"Slide {slide_number}",
                "source_type": "pptx",
            })

    print(
        f"PPTX processed successfully. "
        f"Slides/chunks created: {len(chunks)}"
    )

    return chunks


def process_text_file(file_path: str):

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:

        text = file.read()

    chunks = []

    for chunk in split_text(text):

        chunks.append({
            "id": str(uuid.uuid4()),
            "text": chunk,
            "source": os.path.basename(file_path),
            "page": 1,
            "section": None,
            "source_type": Path(file_path).suffix.lower(),
        })

    return chunks


def process_document(file_path: str):

    extension = Path(file_path).suffix.lower()

    if extension == ".pdf":
        return process_pdf(file_path)

    if extension == ".pptx":
        return process_pptx(file_path)

    if extension in [".md", ".txt"]:
        return process_text_file(file_path)

    raise ValueError(
        f"Unsupported document type: {extension}"
    )